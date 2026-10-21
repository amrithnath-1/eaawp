"""
Dataset parsing, typing and loading (pure standard library + openpyxl; no pandas).

Limits keep the work bounded and safe: 5 MB, 50,000 rows, 100 columns. Excel files are read with
cached values only (``data_only``) - formulas are never evaluated, macros never run. Exports neutralise
spreadsheet formula injection.
"""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime
from statistics import median

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLS = 100
_DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%SZ")
_NUM_STRIP = re.compile(r"[,$€£₹\s]")


class DatasetError(ValueError):
    """A problem the uploader can fix; the message is safe to show."""


def _clean_headers(raw: list) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for i, h in enumerate(raw):
        name = re.sub(r"\s+", " ", str(h if h is not None else "")).strip()[:60] or f"column_{i + 1}"
        base, n = name, seen.get(name.lower(), 0)
        seen[name.lower()] = n + 1
        names.append(base if n == 0 else f"{base}_{n + 1}")
    return names


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    raise DatasetError("The file could not be read as text.")


def _rows_from_csv(data: bytes) -> list[list[str]]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _rows_from_xlsx(data: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DatasetError("The Excel file could not be opened (is it a valid .xlsx?).") from exc
    ws = wb.worksheets[0]
    rows: list[list] = []
    for r in ws.iter_rows(values_only=True):
        rows.append(["" if c is None else (c.isoformat() if isinstance(c, (datetime, date)) else c) for c in r])
        if len(rows) > MAX_ROWS + 1:
            break
    wb.close()
    return rows


def parse_upload(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    """-> (headers, rows as strings). Raises DatasetError with a helpful message."""
    if not data:
        raise DatasetError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise DatasetError(f"The file is larger than the {MAX_BYTES // (1024 * 1024)} MB limit.")
    lower = (filename or "").lower()
    if lower.endswith(".xlsx"):
        rows = _rows_from_xlsx(data)
    elif lower.endswith((".csv", ".tsv", ".txt")):
        rows = _rows_from_csv(data)
    else:
        raise DatasetError("Only .csv and .xlsx files are supported.")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise DatasetError("The file needs a header row and at least one data row.")
    headers = _clean_headers(rows[0])
    if len(headers) > MAX_COLS:
        raise DatasetError(f"More than {MAX_COLS} columns.")
    body = rows[1:]
    if len(body) > MAX_ROWS:
        raise DatasetError(f"More than {MAX_ROWS:,} rows. Upload a smaller extract.")
    width = len(headers)
    norm = [[("" if c is None else str(c)).strip() for c in (r + [""] * width)[:width]] for r in body]
    return headers, norm


def parse_number(s: str) -> float | None:
    t = _NUM_STRIP.sub("", s or "")
    if not t:
        return None
    pct = t.endswith("%")
    t = t.rstrip("%")
    if t.startswith("(") and t.endswith(")"):
        t = "-" + t[1:-1]
    try:
        v = float(t)
    except ValueError:
        return None
    if v != v or v in (float("inf"), float("-inf")):
        return None
    return v / 100 if pct and False else v


def parse_date(s: str) -> date | None:
    s = (s or "").strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def infer_columns(headers: list[str], rows: list[list[str]]) -> list[dict]:
    cols = []
    for i, name in enumerate(headers):
        vals = [r[i] for r in rows if r[i] != ""]
        n = len(vals)
        numeric = sum(1 for v in vals if parse_number(v) is not None)
        dates = sum(1 for v in vals if parse_date(v) is not None)
        if n and numeric / n >= 0.9:
            typ = "number"
        elif n and dates / n >= 0.9:
            typ = "date"
        else:
            typ = "text"
        cols.append({"name": name, "type": typ, "non_null": n, "distinct": len(set(vals)),
                     "sample": list(dict.fromkeys(vals))[:3]})
    return cols


def to_csv_text(headers: list[str], rows: list[list[str]]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    return buf.getvalue()


def load_table(content_csv: str, columns: list[dict]) -> list[dict]:
    """Typed rows: numbers -> float|None, dates -> date|None, text -> str."""
    reader = csv.reader(io.StringIO(content_csv))
    next(reader, None)
    types = [c["type"] for c in columns]
    names = [c["name"] for c in columns]
    out = []
    for r in reader:
        row = {}
        for i, name in enumerate(names):
            v = r[i] if i < len(r) else ""
            row[name] = parse_number(v) if types[i] == "number" else parse_date(v) if types[i] == "date" else v
        out.append(row)
    return out


def safe_cell(v) -> str:
    """Neutralise spreadsheet formula injection in exports."""
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") and not re.fullmatch(r"-?\d+(\.\d+)?", s) else s


def export_csv(columns: list[str], rows: list[list]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([safe_cell(c) for c in columns])
    for r in rows:
        w.writerow([safe_cell(c) for c in r])
    return buf.getvalue()
